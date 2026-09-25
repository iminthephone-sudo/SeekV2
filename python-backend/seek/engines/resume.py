"""Resume builder: render a profile to HTML / Markdown / plain text / DOCX and coach it.

The HTML is deliberately limited to the subset Qt's rich-text engine supports
(tables, headings, lists, inline font styles). The shell shows it in a
QTextBrowser and prints the same document to PDF, so what staff preview is
exactly what gets exported.
"""

from __future__ import annotations

import html
from pathlib import Path
from typing import Any

from .nlp import NLP
from .profiles import profile_text

TEMPLATE_STYLES = {
    # name: (font family, accent colour, name size pt, heading size pt, body size pt)
    "classic": ("Georgia, 'Times New Roman', serif", "#1f2937", 22, 11, 10.5),
    "modern": ("'Segoe UI', 'Helvetica Neue', Arial, sans-serif", "#0b6aa2", 24, 11, 10.5),
    "compact": ("Calibri, Arial, sans-serif", "#374151", 18, 10, 9.5),
}

SECTION_TITLES = {
    "summary": "Summary", "skills": "Skills", "experience": "Experience", "certifications": "Certifications",
    "training": "Training & Programs", "education": "Education", "volunteer": "Volunteer & Community",
    "references": "References",
}
SECTION_ORDER = ["summary", "skills", "experience", "certifications", "training", "education", "volunteer", "references"]


def _e(text: Any) -> str:
    return html.escape(str(text or ""), quote=True)


def _dates(entry: dict[str, Any]) -> str:
    start, end = entry.get("start", ""), entry.get("end", "")
    if start and end:
        return f"{start} – {end}"
    return start or end or entry.get("date", "")


def _contact_line(profile: dict[str, Any]) -> list[str]:
    c = profile.get("contact", {})
    place = ", ".join(filter(None, [c.get("city"), c.get("state")]))
    return [p for p in [place, c.get("phone"), c.get("email"), c.get("linkedin"), c.get("website")] if p]


def _has(profile: dict[str, Any], section: str) -> bool:
    if section == "summary":
        return bool(profile.get("summary"))
    if section == "references":
        opts = profile.get("options", {})
        return bool(profile.get("references") and opts.get("include_references")) or bool(opts.get("references_on_request"))
    return bool(profile.get(section))


# -- HTML ---------------------------------------------------------------------------
def render_html(profile: dict[str, Any], template: str | None = None) -> str:
    template = template or profile.get("options", {}).get("template", "classic")
    font, accent, name_pt, head_pt, body_pt = TEMPLATE_STYLES.get(template, TEMPLATE_STYLES["classic"])
    c = profile.get("contact", {})
    name = c.get("full_name") or profile.get("name", "")
    align = "left" if template == "modern" else "center"
    out = [
        "<html><head><meta charset='utf-8'></head>",
        f"<body style=\"font-family:{font}; font-size:{body_pt}pt; color:#111827;\">",
        f"<p align='{align}' style='margin:0;'><span style='font-size:{name_pt}pt; font-weight:600; color:{accent};'>"
        f"{_e(name)}</span></p>",
    ]
    if profile.get("headline"):
        out.append(f"<p align='{align}' style='margin:2px 0 0 0; font-size:{body_pt + 1}pt; color:#374151;'>"
                   f"{_e(profile['headline'])}</p>")
    contact = _contact_line(profile)
    if contact:
        out.append(f"<p align='{align}' style='margin:4px 0 6px 0; color:#4b5563;'>{' &nbsp;|&nbsp; '.join(_e(x) for x in contact)}</p>")

    def heading(title: str) -> None:
        out.append(f"<h3 style='font-size:{head_pt}pt; color:{accent}; margin:12px 0 2px 0; letter-spacing:1px;'>"
                   f"{_e(title.upper())}</h3><hr style='margin:0 0 4px 0;'/>")

    def row(left: str, right: str) -> None:
        out.append("<table width='100%' cellspacing='0' cellpadding='0' style='margin-top:4px;'><tr>"
                   f"<td>{left}</td><td align='right' style='color:#4b5563;'>{_e(right)}</td></tr></table>")

    def bullets(items: list[str]) -> None:
        if items:
            out.append("<ul style='margin-top:0; margin-bottom:0;'>" +
                       "".join(f"<li>{_e(b)}</li>" for b in items) + "</ul>")

    for section in SECTION_ORDER:
        if not _has(profile, section):
            continue
        heading(SECTION_TITLES[section])
        if section == "summary":
            out.append(f"<p style='margin:2px 0;'>{_e(profile['summary'])}</p>")
        elif section == "skills":
            sep = " &nbsp;•&nbsp; " if template != "compact" else ", "
            out.append(f"<p style='margin:2px 0;'>{sep.join(_e(s) for s in profile['skills'])}</p>")
        elif section == "experience":
            for exp in profile["experience"]:
                left = f"<b>{_e(exp.get('title'))}</b>"
                org = ", ".join(filter(None, [exp.get("employer"), exp.get("location")]))
                if org:
                    left += f" — {_e(org)}"
                row(left, _dates(exp))
                bullets(exp.get("bullets", []))
        elif section == "volunteer":
            for vol in profile["volunteer"]:
                row(f"<b>{_e(vol.get('role'))}</b> — {_e(vol.get('organization'))}", _dates(vol))
                bullets(vol.get("bullets", []))
        elif section == "education":
            for edu in profile["education"]:
                cred = " in ".join(filter(None, [edu.get("credential"), edu.get("field")]))
                school = ", ".join(filter(None, [edu.get("school"), edu.get("location")]))
                row(f"<b>{_e(cred)}</b>" + (f" — {_e(school)}" if school else ""), _dates(edu))
                if edu.get("details"):
                    out.append(f"<p style='margin:0 0 0 12px; color:#374151;'>{_e(edu['details'])}</p>")
        elif section == "certifications":
            for cert in profile["certifications"]:
                extra = f" — {_e(cert['issuer'])}" if cert.get("issuer") else ""
                when = cert.get("date", "") + (f" (exp. {cert['expires']})" if cert.get("expires") else "")
                row(f"<b>{_e(cert.get('name'))}</b>{extra}", when)
        elif section == "training":
            for tr in profile["training"]:
                extra = f" — {_e(tr['provider'])}" if tr.get("provider") else ""
                hours = f" ({_e(tr['hours'])} hrs)" if tr.get("hours") else ""
                row(f"<b>{_e(tr.get('name'))}</b>{extra}{hours}", tr.get("date", ""))
                if tr.get("details"):
                    out.append(f"<p style='margin:0 0 0 12px; color:#374151;'>{_e(tr['details'])}</p>")
        elif section == "references":
            if profile.get("options", {}).get("include_references") and profile.get("references"):
                for ref in profile["references"]:
                    bits = ", ".join(filter(None, [ref.get("relationship"), ref.get("phone"), ref.get("email")]))
                    out.append(f"<p style='margin:2px 0;'><b>{_e(ref.get('name'))}</b> — {_e(bits)}</p>")
            else:
                out.append("<p style='margin:2px 0;'>Available upon request.</p>")
    out.append("</body></html>")
    return "\n".join(out)


# -- Markdown / text -------------------------------------------------------------------
def render_markdown(profile: dict[str, Any]) -> str:
    c = profile.get("contact", {})
    lines = [f"# {c.get('full_name') or profile.get('name', '')}"]
    if profile.get("headline"):
        lines.append(f"**{profile['headline']}**")
    contact = _contact_line(profile)
    if contact:
        lines.append(" | ".join(contact))
    for section in SECTION_ORDER:
        if not _has(profile, section):
            continue
        lines += ["", f"## {SECTION_TITLES[section]}"]
        if section == "summary":
            lines.append(profile["summary"])
        elif section == "skills":
            lines.append(", ".join(profile["skills"]))
        elif section in ("experience", "volunteer"):
            for e in profile[section]:
                title = e.get("title") or e.get("role")
                org = ", ".join(filter(None, [e.get("employer") or e.get("organization"), e.get("location")]))
                lines.append(f"**{title}** — {org} ({_dates(e)})" if _dates(e) else f"**{title}** — {org}")
                lines += [f"- {b}" for b in e.get("bullets", [])]
        elif section == "education":
            for e in profile["education"]:
                cred = " in ".join(filter(None, [e.get("credential"), e.get("field")]))
                lines.append(f"**{cred}** — {e.get('school', '')} {('(' + _dates(e) + ')') if _dates(e) else ''}".rstrip())
                if e.get("details"):
                    lines.append(f"  {e['details']}")
        elif section == "certifications":
            lines += [f"- {x.get('name')}" + (f" — {x['issuer']}" if x.get("issuer") else "")
                      + (f" ({x['date']})" if x.get("date") else "") for x in profile["certifications"]]
        elif section == "training":
            lines += [f"- {x.get('name')}" + (f" — {x['provider']}" if x.get("provider") else "")
                      + (f" ({x['hours']} hrs)" if x.get("hours") else "") for x in profile["training"]]
        elif section == "references":
            if profile.get("options", {}).get("include_references") and profile.get("references"):
                lines += [f"- {r.get('name')} — " + ", ".join(filter(None, [r.get('relationship'), r.get('phone'), r.get('email')]))
                          for r in profile["references"]]
            else:
                lines.append("Available upon request.")
    return "\n".join(lines).strip() + "\n"


def render_text(profile: dict[str, Any]) -> str:
    md = render_markdown(profile)
    out = []
    for line in md.splitlines():
        if line.startswith("# "):
            out.append(line[2:].upper())
        elif line.startswith("## "):
            title = line[3:].upper()
            out += [title, "-" * len(title)]
        else:
            out.append(line.replace("**", "").replace("- ", "• ", 1) if line.startswith("- ") else line.replace("**", ""))
    return "\n".join(out)


# -- DOCX --------------------------------------------------------------------------------
def export_docx(profile: dict[str, Any], path: str | Path) -> str:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt, RGBColor
    except ImportError as exc:  # pragma: no cover - depends on environment
        raise RuntimeError("DOCX export needs python-docx: pip install python-docx") from exc

    template = profile.get("options", {}).get("template", "classic")
    _, accent, name_pt, head_pt, body_pt = TEMPLATE_STYLES.get(template, TEMPLATE_STYLES["classic"])
    rgb = RGBColor.from_string(accent.lstrip("#").upper())
    doc = Document()
    style = doc.styles["Normal"]
    style.font.size = Pt(body_pt)
    style.font.name = "Calibri" if template != "classic" else "Georgia"
    for section in doc.sections:
        section.left_margin = section.right_margin = Pt(54)
        section.top_margin = section.bottom_margin = Pt(46)

    c = profile.get("contact", {})
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if template == "modern" else WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(c.get("full_name") or profile.get("name", ""))
    run.bold, run.font.size, run.font.color.rgb = True, Pt(name_pt), rgb
    for text in filter(None, [profile.get("headline"), "  |  ".join(_contact_line(profile))]):
        para = doc.add_paragraph(text)
        para.alignment = p.alignment

    def heading(title: str) -> None:
        para = doc.add_paragraph()
        r = para.add_run(title.upper())
        r.bold, r.font.size, r.font.color.rgb = True, Pt(head_pt), rgb
        para.paragraph_format.space_before = Pt(10)
        para.paragraph_format.space_after = Pt(2)

    def line(left: str, right: str = "") -> None:
        para = doc.add_paragraph()
        para.add_run(left).bold = True
        if right:
            para.add_run(f"    {right}")
        para.paragraph_format.space_after = Pt(0)

    for section in SECTION_ORDER:
        if not _has(profile, section):
            continue
        heading(SECTION_TITLES[section])
        if section == "summary":
            doc.add_paragraph(profile["summary"])
        elif section == "skills":
            doc.add_paragraph(" • ".join(profile["skills"]))
        elif section in ("experience", "volunteer"):
            for e in profile[section]:
                title = e.get("title") or e.get("role") or ""
                org = ", ".join(filter(None, [e.get("employer") or e.get("organization"), e.get("location")]))
                line(f"{title} — {org}" if org else title, _dates(e))
                for b in e.get("bullets", []):
                    doc.add_paragraph(b, style="List Bullet")
        elif section == "education":
            for e in profile["education"]:
                cred = " in ".join(filter(None, [e.get("credential"), e.get("field")]))
                line(f"{cred} — {e.get('school', '')}", _dates(e))
                if e.get("details"):
                    doc.add_paragraph(e["details"])
        elif section == "certifications":
            for x in profile["certifications"]:
                line(x.get("name", "") + (f" — {x['issuer']}" if x.get("issuer") else ""), x.get("date", ""))
        elif section == "training":
            for x in profile["training"]:
                line(x.get("name", "") + (f" — {x['provider']}" if x.get("provider") else ""), x.get("date", ""))
                if x.get("details"):
                    doc.add_paragraph(x["details"])
        elif section == "references":
            if profile.get("options", {}).get("include_references") and profile.get("references"):
                for r in profile["references"]:
                    doc.add_paragraph(f"{r.get('name')} — " + ", ".join(filter(None, [r.get('relationship'), r.get('phone'), r.get('email')])))
            else:
                doc.add_paragraph("Available upon request.")
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return str(path)


def export(profile: dict[str, Any], fmt: str, path: str | Path, template: str | None = None) -> str:
    fmt = fmt.lower().lstrip(".")
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    if fmt == "docx":
        return export_docx(profile, path)
    if fmt in ("html", "htm"):
        path.write_text(render_html(profile, template), encoding="utf-8")
    elif fmt in ("md", "markdown"):
        path.write_text(render_markdown(profile), encoding="utf-8")
    elif fmt in ("txt", "text"):
        path.write_text(render_text(profile), encoding="utf-8")
    else:
        raise ValueError(f"unsupported format: {fmt} (PDF is printed by the shell)")
    return str(path)


# -- Coaching ----------------------------------------------------------------------------
def analyze(profile: dict[str, Any], nlp: NLP) -> dict[str, Any]:
    """Score the resume and give concrete, per-bullet feedback."""
    bullet_reviews = []
    for section in ("experience", "volunteer"):
        for entry in profile.get(section, []):
            label = entry.get("title") or entry.get("role") or section
            for bullet in entry.get("bullets", []):
                review = nlp.review_bullet(bullet)
                review["entry"] = label
                review["entry_id"] = entry.get("id")
                bullet_reviews.append(review)

    checks = []

    def check(ok: bool, label: str, tip: str) -> None:
        checks.append({"ok": ok, "label": label, "tip": "" if ok else tip})

    c = profile.get("contact", {})
    check(bool(c.get("full_name")), "Name", "Add the participant's name.")
    check(bool(c.get("phone") or c.get("email")), "Phone or email", "Employers need a way to reach them: add a phone number or email.")
    check(bool(c.get("city")), "City/State", "Add a city and state so local employers see they're nearby.")
    check(len(profile.get("summary", "")) >= 60, "Summary", "Write 2–3 sentences on strengths and the kind of work wanted.")
    check(len(profile.get("skills", [])) >= 6, "6+ skills", "List at least 6 skills; use 'Suggest skills' to pull them from the bullets.")
    check(bool(profile.get("experience")), "Work experience", "Add jobs, program work, or crew assignments.")
    check(bool(bullet_reviews) and all(len(e.get("bullets", [])) >= 2 for e in profile.get("experience", [])),
          "2+ bullets per job", "Give each job at least two bullets describing tasks and results.")
    check(bool(profile.get("certifications") or profile.get("training") or profile.get("education")),
          "Credentials or training", "Add certificates, GED/HSE, vocational or program training.")

    placeholders = "[#]" in profile_text(profile)
    check(not placeholders, "No [#] left", "Replace every [#] from the writing helper with a real number, or delete it.")

    listed = {nlp.key(s) for s in profile.get("skills", [])}
    found = nlp.find_skills(profile_text(profile))
    suggested = [s for s, _ in found.most_common() if nlp.key(s) not in listed]

    bullet_score = sum(r["score"] for r in bullet_reviews) / len(bullet_reviews) if bullet_reviews else 0
    check_score = 100 * sum(1 for x in checks if x["ok"]) / len(checks)
    overall = round(0.55 * check_score + 0.45 * bullet_score)
    return {
        "score": overall,
        "checks": checks,
        "bullets": bullet_reviews,
        "suggested_skills": suggested,
        "word_count": len(profile_text(profile).split()),
    }
